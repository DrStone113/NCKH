"""Apply final verified corrections to the generated THS2025-78 working DOCX."""

from io import BytesIO
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image, ImageDraw, ImageFont


ROOT = Path(__file__).resolve().parents[1]
REPORT = ROOT / "docs" / "report_assets" / "final_working.docx"


def replace(paragraph, old, new):
    if old not in paragraph.text:
        return False
    text = paragraph.text.replace(old, new)
    paragraph.clear()
    paragraph.add_run(text)
    return True


def main():
    doc = Document(REPORT)
    for style_name in ("Normal", "Heading 1", "Heading 2", "Heading 3", "Caption", "TOC 1", "TOC 2", "TOC 3"):
        if style_name in doc.styles:
            style = doc.styles[style_name]
            style.font.name = "Times New Roman"
            style._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "Times New Roman")

    replacements = {
        "Technology transfer means:": "Planned technology transfer means (not yet completed):",
        "chỉ số phản ánh thứ hạng trung bình, MRR": "trung bình nghịch đảo thứ hạng của nguồn liên quan đầu tiên, MRR",
        "Riêng bước tìm kiếm mất khoảng 170 mili giây ở mức giữa và 204 mili giây ở mốc 95% lượt.": "Độ trễ của lần chạy trên bản mã hiện hành chưa được đo lại.",
        "https://firebase.google.com/docs/firestore/security/get-started (truy cập": "https://firebase.google.com/docs/auth; https://firebase.google.com/docs/firestore/security/get-started (truy cập",
    }
    for old, new in replacements.items():
        count = sum(replace(p, old, new) for p in doc.paragraphs)
        if count != 1:
            raise RuntimeError(f"Expected exactly one match for {old!r}, found {count}")

    # The achievement row is intentionally empty; keep the official table together.
    for table in doc.tables[4:8]:
        for row in table.rows:
            props = row._tr.get_or_add_trPr()
            if props.find(qn("w:cantSplit")) is None:
                props.append(OxmlElement("w:cantSplit"))
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.keep_with_next = True
    # Put the second-year form on one page rather than stranding a blank row.
    second_year = next(p for p in doc.paragraphs if p.text.strip() == "* Năm thứ 2:")
    second_year.paragraph_format.page_break_before = True
    # A signature must remain beside its administrative label, not alone on a page.
    for row in doc.tables[8].rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.keep_with_next = True
    doc.tables[8].rows[-1].cells[-1].paragraphs[-1].paragraph_format.keep_with_next = False

    image_part = doc.inline_shapes[1]._inline.graphic.graphicData.pic.blipFill.blip.embed
    part = doc.part.related_parts[image_part]
    image = Image.open(BytesIO(part.blob)).convert("RGB")
    draw = ImageDraw.Draw(image)
    draw.rectangle((250, 615, image.width - 15, image.height - 9), fill="white")
    font = ImageFont.truetype("C:/Windows/Fonts/arial.ttf", 31)
    small = ImageFont.truetype("C:/Windows/Fonts/arial.ttf", 25)
    draw.text((295, 661), "Kết quả thực thi  ->  đối chiếu  ->  phản hồi", font=font, fill="#283947")
    draw.text((475, 708), "Lỗi: không xác nhận đã lưu", font=small, fill="#a35624")
    output = BytesIO()
    image.save(output, format="PNG")
    part._blob = output.getvalue()

    doc.save(REPORT)
    print(f"Polished working DOCX: {REPORT}")


if __name__ == "__main__":
    main()
